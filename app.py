import streamlit as st
import pandas as pd
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter
import io
import json
import requests
import base64
import os
import plotly.express as px

st.set_page_config(page_title="Certificaciones Fibra", layout="wide", page_icon="🛠️")

# --- CONFIGURACIÓN GITHUB ---
try:
    GITHUB_TOKEN = st.secrets.get("GITHUB_TOKEN", "")
except:
    GITHUB_TOKEN = ""

GITHUB_USUARIO = "jaxma77-cell"
GITHUB_REPO = "certificaciones-fibra"
GITHUB_ARCHIVO_DATOS = "datos_bot.json"
GITHUB_ARCHIVO_CONFIG = "config_bot.json"

# --- DATOS INICIALES ---
if 'proyectos' not in st.session_state:
    st.session_state.proyectos = {
        "FIBRAMOL JUNIO Y JULIO": {
            "empresa": "Fibranet", "fecha": "Jul-26",
            "conceptos": pd.DataFrame({
                'Concepto': ['Preparación extremo de cable', 'Preparación Sangrado', 
                           'Instalación Caja de Empalme, CTO', 'Instalación Spliter',
                           'Empalme a fusión hasta 64fo', 'Empalme a fusión a partir de 64fo',
                           'Medida de potencia de 1 fibra en 2a y 3a ventana'],
                'Precio': [12.50, 25.00, 8.00, 3.00, 5.00, 2.50, 2.00]
            }),
            "filas": []
        },
        "Santomera Mayo 2024": {
            "empresa": "Fibranet Tecnologia y Sistemas SLU", "fecha": "May-24",
            "conceptos": pd.DataFrame({
                'Concepto': ['Preparación Extremo de Cable', 'Preparación Sangrado',
                           'Instalación Caja de Empalme, CTO', 'Instalación de Spliter',
                           'Empalme a Fusión hasta 24fo', 'Medida de potencia de 1 fibra en 2a y 3a ventana'],
                'Precio': [18.00, 36.00, 12.00, 3.10, 7.00, 2.50]
            }),
            "filas": []
        }
    }
    st.session_state.proyecto_activo = "FIBRAMOL JUNIO Y JULIO"

if 'pie_pagina' not in st.session_state:
    st.session_state.pie_pagina = "F'BERED INGENIERIA EN REDES DE FIBRA"

# --- FUNCIONES AUXILIARES ---
def obtener_conceptos_validos(proyecto):
    df = proyecto["conceptos"]
    return df[df['Concepto'].notna() & (df['Concepto'].astype(str).str.strip() != '')].reset_index(drop=True)

def adaptar_filas(filas, conceptos_validos):
    conceptos_list = conceptos_validos['Concepto'].tolist()
    return [{'Dia': f.get('Dia', ''), 'Nombre': f.get('Nombre', ''), **{c: f.get(c, 0) for c in conceptos_list}} for f in filas]

def calcular_totales(filas, conceptos_list, precios_list):
    return [sum(float(fila.get(c, 0) or 0) * precios_list[i] for i, c in enumerate(conceptos_list)) for fila in filas]

def formato_euro(valor):
    return f"{valor:,.2f} €".replace(",", "X").replace(".", ",").replace("X", ".")

# --- FUNCIÓN: SINCRONIZAR CONCEPTOS CON GITHUB ---
def sincronizar_config_bot(nombre_proyecto_app, conceptos_list, precios_list):
    if not GITHUB_TOKEN:
        return False, "No hay token"
    try:
        url = f"https://api.github.com/repos/{GITHUB_USUARIO}/{GITHUB_REPO}/contents/{GITHUB_ARCHIVO_CONFIG}"
        headers = {"Authorization": f"token {GITHUB_TOKEN}"}
        mapeo = {"FIBRAMOL JUNIO Y JULIO": "FIBRAMOL", "Santomera Mayo 2024": "SANTOMERA"}
        nombre_github = mapeo.get(nombre_proyecto_app, nombre_proyecto_app)
        
        r = requests.get(url, headers=headers)
        sha = None
        datos = {}
        if r.status_code == 200:
            contenido = base64.b64decode(r.json()["content"]).decode("utf-8")
            sha = r.json()["sha"]
            datos = json.loads(contenido)
        
        datos[nombre_github] = {
            "conceptos": conceptos_list,
            "precios": precios_list,
            "num_conceptos": len(conceptos_list)
        }
        
        contenido_nuevo = json.dumps(datos, indent=2, ensure_ascii=False)
        contenido_b64 = base64.b64encode(contenido_nuevo.encode("utf-8")).decode("utf-8")
        payload = {"message": f"Actualizados conceptos de {nombre_github}", "content": contenido_b64}
        if sha: payload["sha"] = sha
        
        r = requests.put(url, headers=headers, json=payload)
        if r.status_code in (200, 201):
            return True, f"Sincronizados {len(conceptos_list)} conceptos en GitHub"
        return False, f"Error: {r.status_code}"
    except Exception as e:
        return False, f"Error: {str(e)}"

# --- FUNCIÓN: BORRAR EN GITHUB ---
def borrar_en_github(tipo, valor, nombre_proyecto_app):
    if not GITHUB_TOKEN: return False, "No hay token"
    try:
        url = f"https://api.github.com/repos/{GITHUB_USUARIO}/{GITHUB_REPO}/contents/{GITHUB_ARCHIVO_DATOS}"
        headers = {"Authorization": f"token {GITHUB_TOKEN}"}
        r = requests.get(url, headers=headers)
        if r.status_code != 200: return False, f"Error al leer GitHub: {r.status_code}"
        
        contenido = base64.b64decode(r.json()["content"]).decode("utf-8")
        sha = r.json()["sha"]
        datos = json.loads(contenido)
        
        mapeo = {"FIBRAMOL JUNIO Y JULIO": "FIBRAMOL", "Santomera Mayo 2024": "SANTOMERA"}
        nombre_github = None
        if nombre_proyecto_app in datos: nombre_github = nombre_proyecto_app
        elif nombre_proyecto_app in mapeo and mapeo[nombre_proyecto_app] in datos: nombre_github = mapeo[nombre_proyecto_app]
        else:
            for proj in list(datos.keys()):
                if proj.upper() in nombre_proyecto_app.upper() or nombre_proyecto_app.upper() in proj.upper():
                    nombre_github = proj; break
        
        if not nombre_github: return False, f"No se encontró el proyecto en GitHub. Proyectos: {list(datos.keys())}"
        
        filas_originales = datos[nombre_github].get("filas", [])
        if tipo == 'dia':
            filas_filtradas = [f for f in filas_originales if f.get('Dia') != valor]; campo = 'Día'
        elif tipo == 'ubicacion':
            filas_filtradas = [f for f in filas_originales if f.get('Nombre') != valor]; campo = 'Nombre'
        else: return False, "Tipo no válido"
        
        num_borradas = len(filas_originales) - len(filas_filtradas)
        if num_borradas == 0: return False, f"No se encontraron filas con {campo}='{valor}'"
        
        datos[nombre_github]["filas"] = filas_filtradas
        contenido_nuevo = json.dumps(datos, indent=2, ensure_ascii=False)
        contenido_b64 = base64.b64encode(contenido_nuevo.encode("utf-8")).decode("utf-8")
        payload = {"message": f"Borrado {num_borradas} filas por {campo}: {valor}", "content": contenido_b64, "sha": sha}
        r = requests.put(url, headers=headers, json=payload)
        
        if r.status_code in (200, 201): return True, f"Borradas {num_borradas} filas de GitHub (proyecto: {nombre_github})"
        return False, f"Error al guardar: {r.status_code}"
    except Exception as e:
        return False, f"Error: {str(e)}"

# --- CSS MODO OSCURO (RESTAURADO) ---
st.markdown("""
<style>
.stApp { background-color: #0e1117; color: #fafafa; }
.stApp * { color: #fafafa !important; }
.stMetric { background-color: #262730; border: 1px solid #3a3a3a; padding: 15px; border-radius: 10px; }
.stMetric label { color: #fafafa !important; }
.stMetric div[data-testid="stMetricValue"] { color: #4CAF50 !important; }
section[data-testid="stSidebar"] { background-color: #1a1a1a; }
section[data-testid="stSidebar"] * { color: #fafafa !important; }
.stButton button { background-color: #0d6efd; color: white !important; }
.stTabs [data-baseweb="tab"] { background-color: #262730; }
.stTabs [aria-selected="true"] { background-color: #0e1117 !important; border-top: 3px solid #0d6efd !important; }
input, select { background-color: #262730 !important; color: #fafafa !important; border: 1px solid #3a3a3a !important; }
div[data-testid="stDataFrame"] { background-color: #262730 !important; border: 1px solid #3a3a3a !important; }
div[data-testid="stDataFrame"] * { color: #fafafa !important; }
.streamlit-expanderHeader { background-color: #262730 !important; border: 1px solid #3a3a3a !important; color: #fafafa !important; }
.streamlit-expanderHeader * { color: #fafafa !important; }
.stAlert { background-color: #262730 !important; border: 1px solid #3a3a3a !important; color: #fafafa !important; }
.stAlert * { color: #fafafa !important; }
</style>
""", unsafe_allow_html=True)

# --- BARRA LATERAL ---
with st.sidebar:
    st.header("📁 PROYECTOS", divider=True)
    proyecto_seleccionado = st.selectbox("Selecciona proyecto:", options=list(st.session_state.proyectos.keys()), index=list(st.session_state.proyectos.keys()).index(st.session_state.proyecto_activo))
    st.session_state.proyecto_activo = proyecto_seleccionado
    st.divider()
    
    if st.button("➕ Nuevo Proyecto", use_container_width=True):
        nombre = st.text_input("Nombre:", "Nuevo Proyecto")
        if nombre and nombre not in st.session_state.proyectos:
            st.session_state.proyectos[nombre] = {"empresa": "", "fecha": "", "conceptos": pd.DataFrame({"Concepto": [""], "Precio": [0.0]}), "filas": []}
            st.session_state.proyecto_activo = nombre; st.rerun()
    
    if st.button("✏️ Renombrar", use_container_width=True):
        nuevo = st.text_input("Nuevo nombre:", st.session_state.proyecto_activo)
        if nuevo and nuevo != st.session_state.proyecto_activo and nuevo not in st.session_state.proyectos:
            st.session_state.proyectos[nuevo] = st.session_state.proyectos.pop(st.session_state.proyecto_activo)
            st.session_state.proyecto_activo = nuevo; st.rerun()
    
    if st.button("🗑️ Eliminar Proyecto", use_container_width=True):
        if len(st.session_state.proyectos) > 1:
            del st.session_state.proyectos[st.session_state.proyecto_activo]
            st.session_state.proyecto_activo = list(st.session_state.proyectos.keys())[0]; st.rerun()
        else: st.error("No puedes eliminar el último proyecto")
    
    st.divider()
    st.subheader("⚙️ Configuración")
    st.session_state.pie_pagina = st.text_input("Pie de página:", st.session_state.pie_pagina)

# --- LÓGICA PRINCIPAL ---
p = st.session_state.proyectos[st.session_state.proyecto_activo]
conceptos_validos = obtener_conceptos_validos(p)
conceptos_list = conceptos_validos['Concepto'].tolist()
precios_list = conceptos_validos['Precio'].tolist()

if p["filas"]: p["filas"] = adaptar_filas(p["filas"], conceptos_validos)
totales = calcular_totales(p["filas"], conceptos_list, precios_list)
total_general = sum(totales) if totales else 0
p["totales"] = totales; p["total_general"] = total_general

st.title(f"️ {st.session_state.proyecto_activo}")
st.caption(f"**{p['empresa']}** · {p['fecha']}")
st.divider()

col1, col2, col3, col4 = st.columns(4)
col1.metric(" TOTAL", formato_euro(total_general))
col2.metric("📋 FILAS", len(p["filas"]))
col3.metric("🏷️ CONCEPTOS", len(conceptos_list))
col4.metric("📊 MEDIA", formato_euro(total_general/len(p["filas"]) if p["filas"] else 0))
st.divider()

tab1, tab2, tab3, tab4 = st.tabs(["📝 REGISTRO", "📊 ANÁLISIS", "📥 IMPORTAR", "⚙️ CONFIGURACIÓN"])

with tab1:
    if not conceptos_list: st.error("❌ Añade conceptos en CONFIGURACIÓN")
    else:
        c1, c2, c3 = st.columns(3)
        if c1.button("➕ AÑADIR FILA", use_container_width=True):
            p["filas"].append({'Dia': '', 'Nombre': '', **{c: 0 for c in conceptos_list}}); st.rerun()
        if c2.button("📋 DUPLICAR ÚLTIMA", use_container_width=True, disabled=not p["filas"]):
            if p["filas"]:
                copia = p["filas"][-1].copy(); copia['Nombre'] = ''; p["filas"].append(copia); st.rerun()
        if c3.button("🗑️ BORRAR TODO", use_container_width=True):
            p["filas"] = []; st.rerun()
        
        st.divider()
        
        # SECCIÓN DE ELIMINAR - MUY VISIBLE
        st.markdown("### 🗑️ ELIMINAR REGISTROS DE GITHUB")
        st.info("⚠️ **Atención**: Al borrar aquí, también se eliminarán del archivo del bot en GitHub.")
        st.caption(f"📌 Proyecto activo: '{st.session_state.proyecto_activo}' | Token: {'✅' if GITHUB_TOKEN else '❌'}")
        
        col_elim1, col_elim2 = st.columns(2)
        
        with col_elim1:
            if p["filas"]:
                dias = sorted(set(f.get('Dia', '') for f in p["filas"] if f.get('Dia')))
                if dias:
                    dia_a_eliminar = st.selectbox("📅 Eliminar todo el día:", [""] + dias, key="sel_elim_dia")
                    if dia_a_eliminar:
                        num_filas_dia = len([f for f in p["filas"] if f.get('Dia') == dia_a_eliminar])
                        if st.button(f"🗑️ Borrar {num_filas_dia} filas del día '{dia_a_eliminar}'", use_container_width=True):
                            p["filas"] = [f for f in p["filas"] if f.get('Dia') != dia_a_eliminar]
                            ok, msg = borrar_en_github('dia', dia_a_eliminar, st.session_state.proyecto_activo)
                            if ok: st.success(f"✅ {msg}")
                            else: st.warning(f"⚠️ Borrado localmente, pero: {msg}")
                            st.rerun()
                else:
                    st.write("No hay días registrados")
            else:
                st.write("No hay filas para eliminar")
        
        with col_elim2:
            if p["filas"]:
                nombres = sorted(set(f.get('Nombre', '') for f in p["filas"] if f.get('Nombre')))
                if nombres:
                    nombre_a_eliminar = st.selectbox("📍 Eliminar por nombre:", [""] + nombres, key="sel_elim_nombre")
                    if nombre_a_eliminar:
                        num_filas_nombre = len([f for f in p["filas"] if f.get('Nombre') == nombre_a_eliminar])
                        if st.button(f"🗑️ Borrar {num_filas_nombre} filas de '{nombre_a_eliminar}'", use_container_width=True):
                            p["filas"] = [f for f in p["filas"] if f.get('Nombre') != nombre_a_eliminar]
                            ok, msg = borrar_en_github('ubicacion', nombre_a_eliminar, st.session_state.proyecto_activo)
                            if ok: st.success(f"✅ {msg}")
                            else: st.warning(f"⚠️ Borrado localmente, pero: {msg}")
                            st.rerun()
                else:
                    st.write("No hay nombres registrados")
            else:
                st.write("No hay filas para eliminar")

        st.divider()
        col_f1, col_f2 = st.columns(2)
        filtro_nombre = col_f1.text_input("🔍 BUSCAR:", placeholder="ALT-CE07")
        dias = sorted(set(f.get('Dia', '') for f in p["filas"] if f.get('Dia')))
        filtro_dia = col_f2.selectbox("📅 FILTRAR POR DÍA:", ["TODOS"] + dias)
        
        filas_vistas = p["filas"]
        if filtro_nombre: filas_vistas = [f for f in filas_vistas if filtro_nombre.upper() in str(f.get('Nombre', '')).upper()]
        if filtro_dia != "TODOS": filas_vistas = [f for f in filas_vistas if f.get('Dia') == filtro_dia]
        st.caption(f"Viendo {len(filas_vistas)} de {len(p['filas'])} filas")
        
        if p["filas"]:
            df = pd.DataFrame(filas_vistas)
            for c in conceptos_list:
                if c not in df.columns: df[c] = 0
            df = df[['Dia', 'Nombre'] + conceptos_list]
            col_config = {"Dia": st.column_config.TextColumn("DÍA", width="small"), "Nombre": st.column_config.TextColumn("NOMBRE")}
            for c in conceptos_list: col_config[c] = st.column_config.NumberColumn(c, min_value=0, step=1)
            
            with st.form("editar"):
                df_edit = st.data_editor(df, num_rows="dynamic", column_config=col_config, hide_index=True, use_container_width=True)
                if st.form_submit_button("💾 GUARDAR CAMBIOS", use_container_width=True, type="primary"):
                    nuevas = df_edit.fillna(0).to_dict('records')
                    for fila in nuevas:
                        fila["total"] = sum(float(fila.get(c, 0)) * precios_list[i] for i, c in enumerate(conceptos_list))
                    if filtro_nombre or filtro_dia != "TODOS":
                        no_filtradas = [f for f in p["filas"] if f not in filas_vistas]
                        p["filas"] = no_filtradas + nuevas
                    else: p["filas"] = nuevas
                    st.success(f"✅ {len(nuevas)} filas guardadas"); st.rerun()
        
        if p["filas"]:
            st.divider()
            st.markdown("#### 📍 RESUMEN POR NOMBRE")
            resumen = {}
            for i, f in enumerate(p["filas"]):
                nom = f.get('Nombre', '') or '(sin nombre)'
                resumen[nom] = resumen.get(nom, {'filas': 0, 'total': 0})
                resumen[nom]['filas'] += 1; resumen[nom]['total'] += totales[i]
            st.dataframe(pd.DataFrame([{'Nombre': n, 'Registros': d['filas'], 'Total': formato_euro(d['total'])} for n, d in sorted(resumen.items(), key=lambda x: x[1]['total'], reverse=True)]), use_container_width=True, hide_index=True)

with tab2:
    if not p["filas"]: st.info("Sin datos")
    else:
        st.markdown("### 📈 TOTAL POR CONCEPTO")
        totales_concepto = [sum(float(f.get(c, 0) or 0) * precios_list[i] for f in p["filas"]) for i, c in enumerate(conceptos_list)]
        fig = px.bar(x=conceptos_list, y=totales_concepto, labels={'x': 'Concepto', 'y': 'Total (€)'}, color=totales_concepto, color_continuous_scale='Blues')
        fig.update_layout(height=400, showlegend=False, xaxis_tickangle=-30, plot_bgcolor='#0e1117', paper_bgcolor='#0e1117')
        st.plotly_chart(fig, use_container_width=True)
        st.divider()
        st.markdown("### 📋 ESTADÍSTICAS")
        stats = []
        for i, c in enumerate(conceptos_list):
            cantidades = [float(f.get(c, 0) or 0) for f in p["filas"] if float(f.get(c, 0) or 0) > 0]
            stats.append({'Concepto': c, 'Precio': formato_euro(precios_list[i]), 'Veces usado': len(cantidades), 'Unidades tot.': int(sum(cantidades)), 'Total': formato_euro(totales_concepto[i])})
        st.dataframe(pd.DataFrame(stats), use_container_width=True, hide_index=True)

with tab3:
    tab_tel, tab_exc = st.tabs(["🤖 DESDE TELEGRAM", "📄 DESDE EXCEL"])
    with tab_tel:
        st.info("Importar datos guardados por el bot de Telegram en GitHub")
        if not GITHUB_TOKEN: st.warning("⚠️ Configura GITHUB_TOKEN en Secrets")
        else:
            if st.button("🔄 LEER DATOS DEL BOT", use_container_width=True, type="primary"):
                try:
                    url = f"https://api.github.com/repos/{GITHUB_USUARIO}/{GITHUB_REPO}/contents/{GITHUB_ARCHIVO_DATOS}"
                    headers = {"Authorization": f"token {GITHUB_TOKEN}"}
                    r = requests.get(url, headers=headers)
                    if r.status_code == 200:
                        contenido = base64.b64decode(r.json()["content"]).decode("utf-8")
                        datos_bot = json.loads(contenido)
                        st.success("✅ Datos leídos correctamente")
                        if datos_bot:
                            st.markdown("**Proyectos disponibles en el bot:**")
                            for pn, dp in datos_bot.items():
                                nf = len(dp.get("filas", [])); tp = sum(f.get("total", 0) for f in dp.get("filas", []))
                                st.markdown(f"- **{pn}**: {nf} filas → {formato_euro(tp)}")
                            st.session_state.datos_bot = datos_bot; st.rerun()
                    elif r.status_code == 404: st.warning("️ El archivo no existe aún. Envía datos al bot primero.")
                except Exception as e: st.error(f"❌ Error: {e}")
            
            if 'datos_bot' in st.session_state and st.session_state.datos_bot:
                st.divider()
                proy_sel = st.selectbox("¿Qué proyecto del bot importar?", list(st.session_state.datos_bot.keys()))
                filas_bot = st.session_state.datos_bot[proy_sel].get("filas", [])
                if filas_bot:
                    st.markdown(f"**{len(filas_bot)} filas disponibles para importar**")
                    modo = st.radio("¿Cómo importar?", ["Añadir a las existentes", "Reemplazar todas"], horizontal=True)
                    if st.button("🚀 IMPORTAR DATOS", type="primary", use_container_width=True):
                        nuevas = []
                        for f in filas_bot:
                            fila = {'Dia': '', 'Nombre': f['Nombre']}
                            cantidades = f.get('cantidades', [])
                            for i, c in enumerate(conceptos_list): fila[c] = cantidades[i] if i < len(cantidades) else 0
                            nuevas.append(fila)
                        if modo == "Añadir a las existentes": p["filas"].extend(nuevas)
                        else: p["filas"] = nuevas
                        st.success(f"✅ ¡{len(nuevas)} filas importadas!"); st.balloons(); st.rerun()
    
    with tab_exc:
        st.info("Importar desde un Excel generado previamente")
        archivo = st.file_uploader("Selecciona archivo .xlsx", type=["xlsx"])
        if archivo:
            try:
                wb = openpyxl.load_workbook(archivo); ws = wb.active
                empresa = ws['A1'].value or ""; nombre_proy = ws['A2'].value or "Importado"; fecha = ws['A3'].value or ""
                conceptos_imp, precios_imp = [], []; col = 3
                while ws.cell(row=4, column=col).value and str(ws.cell(row=4, column=col).value).strip().upper() != 'TOTAL':
                    conceptos_imp.append(ws.cell(row=4, column=col).value); precios_imp.append(float(ws.cell(row=5, column=col).value or 0)); col += 1
                filas_imp = []; row = 6
                while ws.cell(row=row, column=1).value or ws.cell(row=row, column=2).value:
                    fila = {'Dia': ws.cell(row=row, column=1).value or '', 'Nombre': ws.cell(row=row, column=2).value or ''}
                    for i, c in enumerate(conceptos_imp):
                        val = ws.cell(row=row, column=i+3).value; fila[c] = int(float(val)) if val else 0
                    filas_imp.append(fila); row += 1
                st.success(f"✅ Leído: {len(filas_imp)} filas, {len(conceptos_imp)} conceptos")
                if st.button("💾 CARGAR COMO NUEVO PROYECTO", type="primary", use_container_width=True):
                    nombre_nuevo = f"{nombre_proy} (importado)"
                    st.session_state.proyectos[nombre_nuevo] = {"empresa": empresa, "fecha": str(fecha), "conceptos": pd.DataFrame({"Concepto": conceptos_imp, "Precio": precios_imp}), "filas": filas_imp}
                    st.session_state.proyecto_activo = nombre_nuevo; st.success("✅ Proyecto cargado correctamente"); st.rerun()
            except Exception as e: st.error(f"Error al leer: {e}")

with tab4:
    st.markdown("### 🏷️ CONCEPTOS Y PRECIOS")
    col1, col2 = st.columns(2)
    p["empresa"] = col1.text_input("EMPRESA:", p["empresa"]); p["fecha"] = col2.text_input("FECHA/MES:", p["fecha"])
    st.divider()
    st.write("**Edita conceptos y precios:**")
    edited = st.data_editor(p["conceptos"], num_rows="dynamic", use_container_width=True, column_config={"Concepto": st.column_config.TextColumn(width="large"), "Precio": st.column_config.NumberColumn(format="%.2f", min_value=0.0, step=0.5)})
    p["conceptos"] = edited.reset_index(drop=True)
    
    if st.button(" APLICAR CAMBIOS Y SINCRONIZAR CON EL BOT", use_container_width=True, type="primary"):
        p["filas"] = adaptar_filas(p["filas"], obtener_conceptos_validos(p))
        ok, msg = sincronizar_config_bot(st.session_state.proyecto_activo, conceptos_list, precios_list)
        if ok: st.success(f"✅ Conceptos actualizados localmente. {msg}")
        else: st.warning(f"⚠️ Conceptos actualizados localmente, pero: {msg}")
        st.rerun()
    
    st.divider()
    st.markdown("### 💾 GUARDAR/CARGAR CONFIGURACIÓN")
    col_a, col_b = st.columns(2)
    with col_a:
        config_data = {"proyecto": st.session_state.proyecto_activo, "empresa": p["empresa"], "fecha": p["fecha"], "conceptos": conceptos_validos.to_dict(orient="records")}
        st.download_button("⬇️ DESCARGAR CONFIG (.json)", json.dumps(config_data, indent=2, ensure_ascii=False), f"config_{st.session_state.proyecto_activo.replace(' ', '_')}.json", use_container_width=True)
    with col_b:
        uploaded = st.file_uploader("CARGAR CONFIG (.json)", type=["json"])
        if uploaded:
            try:
                data = json.load(uploaded); p["empresa"] = data.get("empresa", p["empresa"]); p["fecha"] = data.get("fecha", p["fecha"]); p["conceptos"] = pd.DataFrame(data.get("conceptos", [])); p["filas"] = []
                st.success("✅ Configuración cargada"); st.rerun()
            except Exception as e: st.error(f"Error: {e}")
    
    st.divider()
    st.markdown("### 📤 EXPORTAR A EXCEL")
    if st.button("🚀 GENERAR EXCEL", type="primary", use_container_width=True, disabled=not p["filas"]):
        wb = openpyxl.Workbook(); ws = wb.active; ws.title = st.session_state.proyecto_activo[:31]
        num_cols = 2 + len(conceptos_list) + 1; last_col = get_column_letter(num_cols)
        for r in range(1, 4): ws.merge_cells(f'A{r}:{last_col}{r}')
        ws['A1'], ws['A1'].font = p["empresa"], Font(bold=True, size=14); ws['A1'].alignment = Alignment(horizontal="center")
        ws['A2'], ws['A2'].font = st.session_state.proyecto_activo, Font(bold=True, size=12); ws['A2'].alignment = Alignment(horizontal="center")
        ws['A3'] = p["fecha"]; ws['A3'].alignment = Alignment(horizontal="center")
        ws['A4'], ws['B4'] = 'DÍA', 'NOMBRE'
        for i, c in enumerate(conceptos_list): ws.cell(row=4, column=i+3, value=c)
        ws.cell(row=4, column=num_cols, value='TOTAL')
        header_font = Font(bold=True, color="FFFFFF"); header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
        for col in range(1, num_cols + 1):
            cell = ws.cell(row=4, column=col); cell.font = header_font; cell.fill = header_fill; cell.alignment = Alignment(horizontal="center")
        for i, precio in enumerate(precios_list):
            cell = ws.cell(row=5, column=i+3, value=precio); cell.number_format = '#,##0.00 "€"'; cell.alignment = Alignment(horizontal="center")
        start_row = 6
        for idx, fila in enumerate(p["filas"]):
            ws.cell(row=start_row, column=1, value=fila.get('Dia', '')); ws.cell(row=start_row, column=2, value=fila.get('Nombre', ''))
            for i, c in enumerate(conceptos_list):
                val = int(fila.get(c, 0)) if fila.get(c, 0) > 0 else ""; ws.cell(row=start_row, column=i+3, value=val)
            total_cell = ws.cell(row=start_row, column=num_cols, value=p["totales"][idx]); total_cell.number_format = '#,##0.00 "€"'; total_cell.font = Font(bold=True)
            start_row += 1
        ws.merge_cells(f'A{start_row}:B{start_row}'); ws.cell(row=start_row, column=1, value='TOTAL GENERAL'); ws.cell(row=start_row, column=1).font = Font(bold=True, size=12); ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="right")
        gt = ws.cell(row=start_row, column=num_cols, value=total_general); gt.number_format = '#,##0.00 "€"'; gt.font = Font(bold=True, size=12, color="FF0000")
        start_row += 2; ws.merge_cells(f'A{start_row}:{last_col}{start_row}'); ws.cell(row=start_row, column=1, value=st.session_state.pie_pagina); ws.cell(row=start_row, column=1).font = Font(italic=True, size=10); ws.cell(row=start_row, column=1).alignment = Alignment(horizontal="center")
        ws.column_dimensions['A'].width = 12; ws.column_dimensions['B'].width = 20
        for i in range(3, num_cols + 1): ws.column_dimensions[get_column_letter(i)].width = 18
        buffer = io.BytesIO(); wb.save(buffer); buffer.seek(0)
        st.download_button("📥 DESCARGAR EXCEL", buffer, f"Certificacion_{st.session_state.proyecto_activo.replace(' ', '_')}.xlsx", mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", use_container_width=True)
        st.success("✅ Excel generado correctamente")
